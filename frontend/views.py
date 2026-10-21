from __future__ import annotations

import html
from io import BytesIO

import altair as alt
import pandas as pd
import streamlit as st
from docx import Document

from backend.config import BLOCKED_INDICATORS, SOURCE_LABELS
from backend.data_loader import DataBundle
from backend.transformations import (
    ACTUARIAL_DEFINITIONS,
    ARIMA_COLORS,
    aggregate_actuarial_flows,
    actuarial_2026_indicator_catalog,
    actuarial_chart,
    actuarial_kpis,
    brl,
    contribution_rate_catalog,
    demographic_summary,
    executive_selection_summary,
    fiscal_annual_cost_breakdown,
    fiscal_dtp_fractionation,
    fiscal_insufficiency_scenario,
    fiscal_management_commitment,
    fiscal_management_breakdown,
    fiscal_long,
    fiscal_composition_catalog,
    fiscal_financing,
    fiscal_personnel_composition,
    fiscal_projection_scenarios,
    fiscal_pressure_simulation,
    fiscal_simulator_defaults,
    fiscal_pressure_catalog,
    fiscal_rgf_reconciliation,
    official_dtp_rcl_series,
    flow_compatibility,
    flow_multiplan_long,
    FLOW_COLORS,
    IPCA_FACTOR,
    integer,
    plan_executive_snapshots,
    percent,
)
from frontend.components import (
    actuarial_metric_card,
    compact_value,
    fiscal_cost_table,
    fiscal_fraction_card,
    fiscal_projection_summary,
    format_table,
    indicator_card,
    metric_card,
    notice,
    page_header,
    plan_meta,
    population_card,
    release_cards,
    source_strip,
)


MONEY_INDICATORS = [
    "Resultado atuarial",
    "Provisões matemáticas",
    "Ativos garantidores",
    "Folha mensal de ativos",
    "Folha mensal de aposentados",
    "Folha mensal de pensionistas",
]


def _line_chart(frame: pd.DataFrame, x: str, y: str, color: str, tooltip: list[str], y_title: str = "Valor") -> alt.Chart:
    return (
        alt.Chart(frame)
        .mark_line(point=True, strokeWidth=3)
        .encode(
            x=alt.X(f"{x}:O", title=x),
            y=alt.Y(f"{y}:Q", title=y_title),
            color=alt.Color(f"{color}:N", title=color, scale=alt.Scale(range=["#1B4F72", "#D6A84B", "#3C91A8", "#A63D40"])),
            tooltip=tooltip,
        )
        .properties(height=330)
    )


def _render_line_chart(frame: pd.DataFrame, x: str, y: str, color: str, tooltip: list[str], y_title: str = "Valor") -> None:
    """Renderiza apenas séries com dados, evitando avisos de extensão infinita."""
    valid = frame.dropna(subset=[x, y, color]) if not frame.empty else frame
    if valid.empty:
        notice("Não há série publicada para este recorte; a ausência não foi convertida em zero.", "warn")
        return
    st.altair_chart(_line_chart(valid, x, y, color, tooltip, y_title), width="stretch")


FISCAL_MANUAL_SECTIONS = (
    (
        "Objetivo e escopo",
        "A tela separa a série oficial DTP/RCL das aproximações gerenciais. "
        "Nenhum subtotal condicionado substitui a DTP oficial, e 2026 permanece parcial e não anualizado.",
    ),
    (
        "Série oficial",
        "DTP e RCL ajustada usam a mesma competência do RGF. A razão oficial é DTP ÷ RCL ajustada × 100. "
        "Rubricas do RREO não são somadas automaticamente ao RGF.",
    ),
    (
        "Decomposição anual",
        "As folhas civis e militares são anualizadas por folha mensal × 13. A contribuição do ente civil "
        "usa a razão entre somas do fluxo atuarial aplicada à base contributiva. A contribuição patronal militar "
        "continua N/A e nunca é substituída pelos 10,5% dos segurados.",
    ),
    (
        "Insuficiência por plano",
        "Para cada plano e ano, a insuficiência é máx(0; despesas previdenciárias originais − receitas previdenciárias originais). "
        "Os resultados dos planos são somados sem compensação cruzada. O valor é cenário gerencial condicionado, não aporte oficial.",
    ),
    (
        "Comprometimento gerencial da RCL",
        "O numerador soma o subtotal ativo conhecido e a insuficiência por plano. O denominador é a RCL ajustada. "
        "Como a contribuição patronal militar está ausente e não há conciliação contábil com a DTP, o percentual não é indicador oficial da LRF.",
    ),
    (
        "Simulador 2027 a 2035",
        "Na Opção A, o crescimento salarial real incide somente sobre a folha ativa civil e militar. "
        "A contribuição civil varia mecanicamente com sua base. Benefícios, receitas e despesas dos fluxos atuariais "
        "permanecem exatamente como publicados no DRAA 2026. A RCL usa a taxa escolhida sem aplicação adicional de IPCA.",
    ),
    (
        "Fontes e limitações",
        "Fontes: Siconfi RGF/RREO, Avaliação Atuarial 2026, fluxos DRAA 2026 e premissas públicas da projeção fiscal. "
        "Continuam bloqueados o fechamento fiscal de 2026, a patronal militar, a equivalência entre o agregado gerencial e a DTP e a qualificação do valor como aporte oficial.",
    ),
)


def _fiscal_manual_docx() -> bytes:
    document = Document()
    document.add_heading("Manual Técnico do Impacto Fiscal", level=0)
    document.add_paragraph(
        "Este manual explica os indicadores, as fórmulas e os limites da tela Impacto Fiscal do Painel Executivo de Transição Previdenciária do Maranhão."
    )
    for title, body in FISCAL_MANUAL_SECTIONS:
        document.add_heading(title, level=1)
        document.add_paragraph(body)
    document.add_heading("Fórmulas principais", level=1)
    for formula in (
        "DTP/RCL oficial (%) = DTP oficial ÷ RCL ajustada × 100",
        "Insuficiência do plano = máx(0; despesas originais − receitas originais)",
        "Total com insuficiência = subtotal ativo conhecido + soma das insuficiências por plano",
        "Comprometimento gerencial (%) = total com insuficiência ÷ RCL ajustada × 100",
    ):
        document.add_paragraph(formula, style="List Bullet")
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _indicator_manual_docx(plan: str, catalog: pd.DataFrame, manual_code: str) -> bytes:
    """Gera o manual Word do indicador e plano atualmente selecionados."""
    manual_row = catalog.loc[catalog["Sigla"].eq(manual_code)].iloc[0]
    document = Document()
    document.add_heading("Manual Técnico dos Indicadores", level=0)
    document.add_paragraph(f"Plano / massa: {plan}")
    document.add_paragraph(f"Indicador: {manual_row['Sigla']} · {manual_row['Indicador']}")
    fields = (
        ("O que mede", "Definição"),
        ("Objetivo", "Objetivo"),
        ("Fórmula", "Fórmula"),
        ("Componentes", "Componentes"),
        ("Interpretação", "Interpretação"),
        ("Critério", "Critério"),
        ("Limitação", "Limitação"),
        ("Fonte", "Fonte"),
        ("Competência", "Competência"),
    )
    for title, column in fields:
        document.add_heading(title, level=1)
        document.add_paragraph(str(manual_row[column]))
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


@st.dialog("Manual Técnico do Impacto Fiscal", width="large")
def _render_fiscal_manual() -> None:
    st.markdown(
        '<div class="manual-intro">Leitura institucional das fórmulas, fontes, competências e limitações da tela atual.</div>',
        unsafe_allow_html=True,
    )
    for title, body in FISCAL_MANUAL_SECTIONS:
        st.markdown(
            f'<section class="manual-section"><h3>{title}</h3><p>{body}</p></section>',
            unsafe_allow_html=True,
        )
    st.markdown(
        '<div class="manual-formula"><small>Fórmula de comprometimento gerencial</small>'
        'Total com insuficiência ÷ RCL ajustada × 100</div>',
        unsafe_allow_html=True,
    )
    action_columns = st.columns(2)
    with action_columns[0]:
        st.download_button(
            "Exportar manual para Word",
            data=_fiscal_manual_docx(),
            file_name="Manual_Tecnico_Impacto_Fiscal_MA.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            key="download_fiscal_manual",
            width="stretch",
        )
    with action_columns[1]:
        if st.button("Fechar manual", key="close_fiscal_manual", width="stretch"):
            st.rerun()


@st.dialog("Manual Técnico dos Indicadores", width="large")
def _render_indicator_manual(catalog: pd.DataFrame) -> None:
    """Apresenta o catálogo técnico atuarial sem alterar rota, plano ou seção."""
    current_plan = st.session_state.get("diagnosis_plan", "Plano não selecionado")
    st.markdown(
        '<div class="manual-intro">Consulta institucional das definições, fórmulas, componentes, '
        f'critérios e limitações do plano/massa <strong>{html.escape(str(current_plan))}</strong>.</div>',
        unsafe_allow_html=True,
    )
    manual_code = st.selectbox(
        "Indicador do manual",
        list(catalog["Sigla"]),
        key="diagnosis_manual_indicator",
    )
    manual_row = catalog.loc[catalog["Sigla"].eq(manual_code)].iloc[0]
    st.markdown(
        f'<section class="manual-section"><h3>{html.escape(str(manual_row["Indicador"]))}</h3>'
        f'<p><strong>O que mede:</strong> {html.escape(str(manual_row["Definição"]))}</p>'
        f'<p><strong>Objetivo:</strong> {html.escape(str(manual_row["Objetivo"]))}</p>'
        f'<p><strong>Componentes:</strong> {html.escape(str(manual_row["Componentes"]))}</p>'
        f'<p><strong>Interpretação:</strong> {html.escape(str(manual_row["Interpretação"]))}</p>'
        f'<p><strong>Critério:</strong> {html.escape(str(manual_row["Critério"]))}</p>'
        f'<p><strong>Limitação:</strong> {html.escape(str(manual_row["Limitação"]))}</p>'
        f'<p><strong>Fonte e competência:</strong> {html.escape(str(manual_row["Fonte"]))} · '
        f'{html.escape(str(manual_row["Competência"]))}</p></section>',
        unsafe_allow_html=True,
    )
    st.markdown(
        '<div class="manual-formula"><small>Fórmula do indicador selecionado</small>'
        f'{html.escape(str(manual_row["Fórmula"]))}</div>',
        unsafe_allow_html=True,
    )
    action_columns = st.columns(2)
    with action_columns[0]:
        safe_plan = str(current_plan).replace(" / ", "_").replace(" ", "_")
        st.download_button(
            "Exportar indicador para Word",
            data=_indicator_manual_docx(str(current_plan), catalog, str(manual_code)),
            file_name=f"Manual_Tecnico_{manual_code}_{safe_plan}.docx",
            mime="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
            key="download_indicator_manual",
            width="stretch",
        )
    with action_columns[1]:
        if st.button("Fechar manual dos indicadores", key="close_indicator_manual", width="stretch"):
            st.rerun()


def render_executive(bundle: DataBundle) -> None:
    page_header(
        "Visão Executiva",
        "Leitura por plano e massa, com população, folha mensal, síntese atuarial e condicionantes visíveis.",
    )
    snapshots = plan_executive_snapshots(bundle.actuarial)
    selected_plans = st.multiselect(
        "Planos / massas",
        list(snapshots),
        default=["Previdenciário / Civil"],
        max_selections=3,
        key="executive_plans",
    )
    if not selected_plans:
        notice("Selecione de um a três planos. Nenhuma ausência será convertida em zero.", "warn")
        return
    snapshot = executive_selection_summary(bundle.actuarial, bundle.flows, selected_plans)
    selection_label = " + ".join(selected_plans)
    plan_meta(selection_label, snapshot["competence"], snapshot["status"], snapshot["reservation"])

    st.subheader("População e folha mensal")
    population_columns = st.columns(3)
    for column, (group, values) in zip(population_columns, snapshot["populations"].items()):
        with column:
            quantity = integer(values["quantity"]) if values["quantity"] is not None else "N/A"
            payroll = brl(values["payroll"], compact=True) if values["payroll"] is not None else "N/A"
            detail = "31/12/2025 · fonte atuarial" if values["payroll"] is not None else "ausência estrutural ou não publicada; não equivale a zero"
            population_card(group, quantity, payroll, detail)

    notice(
        "Cada quantidade está vinculada à folha da mesma população. Ausência de linha na fonte é exibida como N/A e nunca convertida em zero.",
        "info",
    )

    st.subheader("Alíquota média implícita no fluxo atuarial de 2026")
    implicit = snapshot["implicit_rates"].copy()
    implicit["Alíquota implícita (%)"] = implicit["Alíquota implícita (%)"].map(percent)
    st.dataframe(
        format_table(implicit, money=["Contribuição no fluxo", "Base de cálculo"]),
        width="stretch",
        hide_index=True,
    )
    notice(
        "A alíquota implícita é a razão entre a soma das contribuições e a soma das bases dos planos selecionados. Não é média simples, não substitui alíquota legal e sinais negativos da fonte são normalizados apenas na apresentação percentual.",
        "info",
    )

    st.subheader("Alíquotas previdenciárias vigentes — 2026")
    rates = contribution_rate_catalog(bundle.contribution_rates)
    st.dataframe(rates, width="stretch", hide_index=True)
    notice(
        "A alíquota patronal civil normal é uma fórmula: dobro da contribuição do segurado, com piso de 15% e teto de 44%; não existe percentual único nem taxa média calculada.",
        "info",
    )
    notice(
        "A alíquota patronal suplementar vigente não foi localizada e a contribuição patronal militar autônoma não foi identificada. Ausência normativa não equivale a 0%. Contribuição patronal não é aporte para cobertura de insuficiência financeira.",
        "warn",
    )
    notice("Informação pendente de validação normativa: percentual patronal suplementar e eventual ato específico posterior às fontes consultadas.", "warn")
    st.caption("Contribuição patronal não é aporte para cobertura de insuficiência financeira.")

    if len(selected_plans) > 1:
        synthesis = "O recorte consolida somente os planos selecionados. Contagens e valores monetários são somados sem a linha Sistema; alíquotas são calculadas pela razão entre somas."
    elif selected_plans[0] == "Previdenciário / Civil":
        synthesis = "O FEPA Civil concentra a maior massa e o maior déficit observado. A leitura executiva deve permanecer associada à data-base e não autoriza inferência sobre custeio ainda não validado."
    elif selected_plans[0] == "Financeiro / Civil":
        synthesis = "O bloco Tesouro Civil reúne aposentados e pensionistas. A ausência de ativos e de folha de ativos é estrutural no recorte publicado e não foi convertida em zero executivo."
    else:
        synthesis = "O recorte militar é uma ponte agregada entre 2022 e o SPSM 2025. Ele permite leitura de posição, mas não continuidade plano a plano nem atribuição causal."
    notice(f"Síntese executiva — {synthesis}", "warn" if selected_plans != ["Previdenciário / Civil"] else "info")

    st.divider()
    st.subheader("Síntese dos planos selecionados")
    selected_quantities = [values["quantity"] for values in snapshot["populations"].values()]
    population_complete = all(value is not None for value in selected_quantities)
    selected_population = sum(selected_quantities) if population_complete else None
    cols = st.columns(4)
    with cols[0]:
        metric_card("Resultado atuarial", brl(snapshot["actuarial_result"], compact=True), "Somente planos selecionados · 31/12/2025")
    with cols[1]:
        metric_card("Provisões matemáticas", brl(snapshot["provisions"], compact=True), "Somente planos selecionados · 31/12/2025")
    with cols[2]:
        metric_card("Ativos garantidores", brl(snapshot["guarantee_assets"], compact=True), "Soma dos valores publicados no recorte")
    with cols[3]:
        population_detail = (
            "Ativos, aposentados e pensionistas selecionados"
            if population_complete
            else "N/A — ao menos uma população selecionada não foi publicada"
        )
        metric_card("População", integer(selected_population), population_detail)

    st.subheader("Resultado atuarial dos planos selecionados")
    selected_source_names = [snapshots[plan]["source_name"] for plan in selected_plans]
    chart_data = actuarial_chart(bundle.actuarial, ["Resultado atuarial"], selected_source_names)
    chart = (
        alt.Chart(chart_data)
        .mark_bar(cornerRadiusTopLeft=4, cornerRadiusTopRight=4)
        .encode(
            x=alt.X("Plano / ponte:N", title=None, axis=alt.Axis(labelAngle=0)),
            y=alt.Y("Valor:Q", title="R$"),
            color=alt.Color("Comparação:N", scale=alt.Scale(range=["#8B7A4A", "#C3A35F", "#17534D"])),
            xOffset="Comparação:N",
            tooltip=["Plano / ponte", "Comparação", alt.Tooltip("Valor:Q", format=",.2f")],
        )
        .properties(height=360)
    )
    st.altair_chart(chart, width="stretch")
    notice("O gráfico usa somente os planos selecionados. A leitura real não equivale a equilíbrio atuarial; a ponte militar não demonstra continuidade plano a plano ou causalidade.", "info")
    source_strip(SOURCE_LABELS["actuarial"])


def render_actuarial(bundle: DataBundle) -> None:
    page_header("Diagnóstico Atuarial", "Posição, indicadores e comparação atuarial com a hierarquia visual da ARIMA Intelligence.")
    plans = list(plan_executive_snapshots(bundle.actuarial))
    plan_choice = st.selectbox("Plano / massa do diagnóstico", plans, key="diagnosis_plan")
    catalog = actuarial_2026_indicator_catalog(bundle.actuarial, bundle.actuarial_balances, plan_choice)
    snapshot = plan_executive_snapshots(bundle.actuarial)[plan_choice]
    coverage = catalog.loc[catalog["Sigla"].eq("ICP"), "Valor"].iloc[0]

    plan_meta(
        plan_choice,
        "Avaliação 2026 · data-base 31/12/2025",
        "DISPONÍVEL COM RESSALVAS",
        "Indicadores pontuais; ausência e denominador zero permanecem N/A.",
    )
    section = st.radio(
        "Navegação interna do diagnóstico",
        ["Posição Atuarial", "Indicadores Atuariais", "Comparação 2022–2025"],
        horizontal=True,
        key="diagnosis_section",
    )
    st.caption("Navegação contextual na mesma página; nenhum detalhe abre nova URL, aba ou janela.")

    if section == "Posição Atuarial":
        st.subheader("Posição atuarial")
        cols = st.columns(4)
        cards = [
            ("Resultado atuarial primário", brl(snapshot["actuarial_result"], compact=True), "Ativos garantidores − provisão matemática", "scale", True),
            ("Provisão matemática", brl(snapshot["provisions"], compact=True), "Obrigação atuarial consolidada", "landmark", False),
            ("Ativos garantidores", brl(snapshot["guarantee_assets"], compact=True), "Patrimônio destinado à cobertura", "wallet", False),
            ("Índice de cobertura", percent(coverage), "ICP · cobertura das obrigações", "trend", False),
        ]
        for column, (label, value, detail, icon, featured) in zip(cols, cards):
            with column:
                actuarial_metric_card(label, value, detail, icon, featured)
        notice("A leitura é pontual e restrita a 2026. Resultado negativo expressa déficit atuarial primário; não é tendência nem ranking.", "warn")
        with st.expander("Controles de consistência da base 2026", expanded=False):
            position_columns = ["Sigla", "Classificação", "Competência", "Limitação"]
            position = catalog.loc[
                catalog["Sigla"].isin(["ICP", "RAP"]), position_columns
            ]
            st.dataframe(position, width="stretch", hide_index=True)
            st.caption("Os componentes a valor presente preservam o resíduo de reconciliação, sem redistribuição artificial.")

    elif section == "Indicadores Atuariais":
        heading, manual = st.columns([2.2, 1])
        with heading:
            st.subheader("Indicadores atuariais · avaliação 2026")
            st.caption("Oito indicadores pontuais do catálogo efetivo ARIMA; detalhes técnicos permanecem recolhidos.")
        with manual:
            if st.button(
                "Manual Técnico dos Indicadores",
                key="open_indicator_manual",
                type="secondary",
                width="stretch",
            ):
                _render_indicator_manual(catalog)

        st.markdown('<div class="indicator-grid-marker" aria-hidden="true"></div>', unsafe_allow_html=True)
        for start in range(0, len(catalog), 3):
            columns = st.columns(3)
            for offset, column in enumerate(columns):
                index = start + offset
                if index >= len(catalog):
                    continue
                row = catalog.iloc[index]
                shown = brl(row["Valor"], compact=True) if row["Unidade"] == "R$" else percent(row["Valor"])
                with column:
                    indicator_card(
                        str(row["Sigla"]),
                        str(row["Indicador"]),
                        shown,
                        str(row["Classificação"]),
                        str(row["Interpretação"]),
                        None,
                    )
                    with st.expander(f"Detalhamento técnico · {row['Sigla']}", expanded=False):
                        detail_labels = (
                            "Definição", "Fórmula", "Componentes", "Fonte", "Competência",
                            "Interpretação", "Limitação", "Critério", "Condição de uso",
                        )
                        for label in detail_labels:
                            st.markdown(f"**{label}:** {row[label]}")
        notice("Comparações externas permanecem ocultas: não há base homogênea e validada para ranking entre RPPS neste pacote.", "info")

    else:
        st.subheader("Comparação do resultado atuarial · 2022 e 2025")
        st.caption("Bloco histórico próprio, separado dos indicadores pontuais da avaliação 2026.")
        comparison = st.radio("Tipo de comparação", ["Nominal", "Real (R$ de dez/2025)"], horizontal=True)
        source_plan = plan_executive_snapshots(bundle.actuarial)[plan_choice]["source_name"]
        subset = bundle.actuarial.loc[
            bundle.actuarial["Plano / ponte"].eq(source_plan)
            & bundle.actuarial["Indicador"].eq("Resultado atuarial")
        ].copy()
        start_column = "Valor original 2022" if comparison == "Nominal" else "2022 atualizado a dez/2025"
        chart_data = subset.melt(
            id_vars=["Plano / ponte", "Unidade"],
            value_vars=[start_column, "Valor observado 2025"],
            var_name="Período",
            value_name="Valor",
        )
        chart_data["Período"] = chart_data["Período"].map({start_column: "2022", "Valor observado 2025": "2025"})
        chart = (
            alt.Chart(chart_data)
            .mark_bar(cornerRadiusTopLeft=4, cornerRadiusTopRight=4)
            .encode(
                x=alt.X("Período:N", title=None),
                y=alt.Y("Valor:Q", title="R$"),
                color=alt.Color("Período:N", scale=alt.Scale(range=["#8B7A4A", "#17534D"]), legend=None),
                tooltip=["Plano / ponte", "Período", alt.Tooltip("Valor:Q", format=",.2f")],
            )
            .properties(height=320)
        )
        st.altair_chart(chart, width="stretch")
        st.dataframe(
            format_table(subset[["Plano / ponte", start_column, "Valor observado 2025", "Variação nominal", "Variação real"]], money=[start_column, "Valor observado 2025", "Variação nominal", "Variação real"]),
            width="stretch",
            hide_index=True,
        )
        notice(f"Fator IPCA aprovado: {IPCA_FACTOR:.10f}. A opção real expressa 2022 em R$ de dezembro de 2025; datas-base 31/12/2022 e 31/12/2025. A comparação de estoques não constitui evolução dos indicadores pontuais nem conclusão causal.", "warn")
        source_strip(f"Avaliação Atuarial 2026 · data-base 31/12/2025 · {SOURCE_LABELS['release']}")


def render_flows(bundle: DataBundle) -> None:
    page_header("Fluxos Atuariais", "Receita, despesa e resultado anual agregados para os planos selecionados.")
    evaluation = st.selectbox("Avaliação", [2026], key="flow_evaluation")
    available = bundle.flows.loc[bundle.flows["Avaliação"].eq(evaluation)]
    dimensions = [name for name in plan_executive_snapshots(bundle.actuarial) if name in set(available["Dimensão"])]
    selected_dimensions = st.multiselect(
        "Planos / massas",
        dimensions,
        default=dimensions,
        max_selections=3,
        key="flow_plans",
    )
    if not selected_dimensions:
        notice("Selecione de um a três planos. Nenhuma ausência será convertida em zero.", "warn")
        source_strip(SOURCE_LABELS["flows"])
        return
    selected = available.loc[available["Dimensão"].isin(selected_dimensions)].copy()
    min_year, max_year = int(selected["Exercício"].min()), int(selected["Exercício"].max())
    years = st.slider("Exercício", min_year, max_year, (min_year, min(max_year, min_year + 29)))
    table, long = aggregate_actuarial_flows(bundle.flows, evaluation, selected_dimensions, years)
    compatibility = flow_compatibility(bundle.flows, evaluation, selected_dimensions)
    notice(str(compatibility["reason"]), "info")
    chart = (
        alt.Chart(long)
        .mark_line(point=True, strokeWidth=3)
        .encode(
            x=alt.X("Ano:O", title="Ano", axis=alt.Axis(labelAngle=0, labelLimit=65)),
            y=alt.Y("Valor:Q", title="R$"),
            color=alt.Color("Série:N", title=None, scale=alt.Scale(domain=list(ARIMA_COLORS), range=list(ARIMA_COLORS.values()))),
            tooltip=["Série", "Ano", alt.Tooltip("Valor:Q", format=",.2f")],
        )
        .properties(height=360)
    )
    st.altair_chart(chart, width="stretch")
    st.dataframe(
        format_table(table, money=["Receita agregada", "Despesa agregada", "Resultado anual agregado"]),
        width="stretch",
        hide_index=True,
    )
    assumptions = selected[["Dimensão", "Taxa de juros (%)", "Método", "Competência", "Fonte"]].drop_duplicates("Dimensão")
    assumptions["Taxa de juros (%)"] = assumptions["Taxa de juros (%)"].map(percent)
    with st.expander("Premissas por plano", expanded=False):
        st.dataframe(assumptions, width="stretch", hide_index=True)
    notice("Taxas de juros distintas foram preservadas como premissas por plano. A agregação soma fluxos financeiros anuais; não recalcula desconto, valor presente ou taxa média.", "info")
    notice("Os fluxos futuros apresentados correspondem a projeções atuariais provenientes dos DRAA e não representam projeções fiscais do Estado.", "risk")
    notice("MIL FIN 2026 é N/A estrutural e não aparece como plano com valor zero.", "risk")
    source_strip(SOURCE_LABELS["flows"])


def render_fiscal(bundle: DataBundle) -> None:
    page_header(
        "Impacto Fiscal",
        "Leitura oficial da DTP e simulação gerencial condicionada da pressão previdenciária, sem misturar competências nem preencher ausências.",
    )
    if st.button("Manual Técnico do Impacto Fiscal", key="open_fiscal_manual", type="secondary"):
        _render_fiscal_manual()
    years = st.multiselect("Ano", [2022, 2023, 2024, 2025, 2026], default=[2022, 2023, 2024, 2025, 2026])
    frame = bundle.fiscal.loc[bundle.fiscal["Ano"].isin(years)].copy()
    if frame.empty:
        notice("Selecione ao menos um exercício.", "info")
        return
    latest = frame.sort_values("Ano").iloc[-1]

    st.subheader("Síntese fiscal")
    cards = st.columns(4)
    with cards[0]:
        metric_card("DTP", brl(latest["DTP"], compact=True), f"{int(latest['Ano'])} · RGF {latest['Competência RGF']}")
    with cards[1]:
        metric_card("RCL ajustada", brl(latest["RCL ajustada (pessoal)"], compact=True), "Denominador do mesmo RGF")
    with cards[2]:
        metric_card("DTP / RCL ajustada", percent(latest["DTP / RCL ajustada"]), "Única razão de pressão liberada")
    with cards[3]:
        metric_card("Resultado capitalização", brl(latest["Resultado fundo em capitalização"], compact=True), f"RREO {latest['Competência RREO']}")
    notice("2026 é parcial: RGF em 12 meses móveis até o 2º quadrimestre e RREO acumulado até o 4º bimestre. Nenhum valor foi anualizado.", "warn")

    st.subheader("Série oficial DTP / RCL ajustada — 2022 a 2026")
    official = official_dtp_rcl_series(frame)
    st.dataframe(
        format_table(
            official,
            money=["DTP", "RCL ajustada (pessoal)"],
            pct=["DTP / RCL ajustada"],
        ),
        width="stretch",
        hide_index=True,
    )
    notice("A série oficial usa exclusivamente o RGF. As rubricas do RREO permanecem separadas; 2026 é parcial e não foi anualizado.", "info")

    st.subheader("Composição da despesa com pessoal")
    composition = fiscal_personnel_composition(frame)
    _render_line_chart(composition, "Ano", "Valor", "Componente", ["Ano", "Componente", alt.Tooltip("Valor:Q", format=",.2f")], "R$")
    with st.expander("Rastreabilidade técnica da composição", expanded=False):
        st.dataframe(fiscal_composition_catalog(), width="stretch", hide_index=True)
    notice("A fonte permite separar aposentadorias/reserva/reformas de pensões, mas não autoriza decompor esses valores entre civis e militares. Ativos civis, militares e demais vínculos permanecem condicionados.", "info")

    st.subheader("Decomposição gerencial do custo anual")
    annual_costs = fiscal_annual_cost_breakdown(bundle.actuarial, bundle.flows)
    fiscal_cost_table(annual_costs)
    notice(
        "A folha anual resulta de folha mensal × 13 e é uma aproximação gerencial, não despesa oficial liquidada. A contribuição civil usa a razão entre somas do fluxo atuarial aplicada à respectiva base contributiva. A contribuição patronal militar permanece N/A; os 10,5% dos segurados não são uma alíquota patronal.",
        "warn",
    )
    notice(
        "CENÁRIO GERENCIAL CONDICIONADO — NÃO REPRESENTA DTP OFICIAL. Os totais destacados preservam a contribuição patronal militar como N/A.",
        "risk",
    )

    insufficiency = fiscal_insufficiency_scenario(bundle.flows)
    with st.expander("Memória do cenário de insuficiência por plano", expanded=False):
        st.dataframe(
            format_table(
                insufficiency["components"],
                money=["Receitas", "Despesas", "Insuficiência gerencial"],
            ),
            width="stretch",
            hide_index=True,
        )
        notice(
            "Cenário DRAA 2026 por plano, sem compensação cruzada. O valor não é aporte oficial e depende de validação atuarial e fiscal/contábil.",
            "warn",
        )

    st.subheader("Comprometimento da RCL com servidores efetivos e insuficiência previdenciária")
    commitment = fiscal_management_commitment(frame, annual_costs)
    commitment_columns = st.columns(4)
    commitment_cards = [
        ("Subtotal ativo conhecido", brl(commitment["active_subtotal"], compact=True), "Patronal militar permanece N/A"),
        ("Insuficiência previdenciária", brl(commitment["insufficiency"], compact=True), "Soma por plano, sem compensação"),
        ("Total com insuficiência", brl(commitment["total_with_insufficiency"], compact=True), "Agregado gerencial condicionado"),
        ("Comprometimento gerencial", percent(commitment["management_commitment"]), "Sobre a RCL ajustada; não é indicador oficial LRF"),
    ]
    for column, (label, value, detail) in zip(commitment_columns, commitment_cards):
        with column:
            metric_card(label, value, detail)
    notice(f"{commitment['status']}. {commitment['limitation']}", "risk")
    st.caption(
        f"Referência oficial preservada em separado: DTP {brl(commitment['official_dtp'])}; "
        f"DTP/RCL ajustada {percent(commitment['official_dtp_rcl'])}. O agregado gerencial não reconcilia nem substitui esses valores."
    )

    st.subheader("Simulador de pressão fiscal — 2027 a 2035")
    defaults = fiscal_simulator_defaults(
        bundle.fiscal_projections,
        bundle.fiscal_projection_assumptions,
    )
    scenario_growth = defaults["scenario_salary_growth_percent"]
    presentation_defaults = {
        "scenario": "A",
        "rcl_growth_percent": 2.0,
        "salary_growth_percent": 1.5,
        "projection_year": 2027,
    }

    def sync_scenario_growth() -> None:
        selected_scenario = st.session_state["fiscal_sim_scenario"]
        st.session_state["fiscal_sim_salary_growth"] = scenario_growth[selected_scenario]

    def restore_base_assumptions() -> None:
        st.session_state["fiscal_sim_scenario"] = presentation_defaults["scenario"]
        st.session_state["fiscal_sim_rcl_growth"] = presentation_defaults["rcl_growth_percent"]
        st.session_state["fiscal_sim_salary_growth"] = presentation_defaults["salary_growth_percent"]
        st.session_state["fiscal_projection_year"] = presentation_defaults["projection_year"]

    if "fiscal_sim_scenario" not in st.session_state:
        st.session_state["fiscal_sim_scenario"] = presentation_defaults["scenario"]
    if "fiscal_sim_rcl_growth" not in st.session_state:
        st.session_state["fiscal_sim_rcl_growth"] = presentation_defaults["rcl_growth_percent"]
    if "fiscal_sim_salary_growth" not in st.session_state:
        st.session_state["fiscal_sim_salary_growth"] = presentation_defaults["salary_growth_percent"]
    if "fiscal_projection_year" not in st.session_state:
        st.session_state["fiscal_projection_year"] = presentation_defaults["projection_year"]

    control_columns = st.columns([1, 1, 1, 0.8])
    with control_columns[0]:
        scenario = st.selectbox(
            "Cenário-base",
            ["A", "B", "C"],
            key="fiscal_sim_scenario",
            on_change=sync_scenario_growth,
            help="A/B/C preservam as premissas homologadas de crescimento real da folha ativa.",
        )
    with control_columns[1]:
        rcl_growth = st.number_input(
            "Crescimento anual da RCL (%)",
            min_value=-20.0,
            max_value=30.0,
            step=0.1,
            format="%.4f",
            key="fiscal_sim_rcl_growth",
            help="Taxa nominal constante. A apresentação pública inicia em 2,00% ao ano.",
        )
    with control_columns[2]:
        salary_growth = st.number_input(
            "Crescimento real anual da folha ativa (%)",
            min_value=-20.0,
            max_value=30.0,
            step=0.1,
            format="%.4f",
            key="fiscal_sim_salary_growth",
            help="Opção A: incide somente na folha ativa civil e militar; valor inicial de apresentação de 1,50%.",
        )
    with control_columns[3]:
        st.write("")
        st.button(
            "Restaurar premissas-base",
            key="restore_fiscal_simulator",
            on_click=restore_base_assumptions,
            width="stretch",
        )

    simulation = fiscal_pressure_simulation(
        annual_costs,
        bundle.flows,
        bundle.fiscal_projections,
        scenario,
        rcl_growth,
        salary_growth,
    )
    projection_year = st.radio(
        "Exercício da projeção",
        list(range(2027, 2036)),
        index=0,
        horizontal=True,
        key="fiscal_projection_year",
        help="Selecione o exercício para atualizar o destaque e a composição executiva abaixo.",
    )
    selected_projection = simulation.loc[simulation["Ano"].eq(projection_year)].iloc[0]
    fiscal_projection_summary(
        int(selected_projection["Ano"]),
        str(selected_projection["Cenário-base"]),
        percent(selected_projection["Comprometimento gerencial"]),
        brl(selected_projection["Subtotal ativo conhecido"], compact=True),
        brl(selected_projection["Insuficiência total"], compact=True),
        brl(selected_projection["Total com insuficiência"], compact=True),
        brl(selected_projection["RCL ajustada projetada"], compact=True),
    )

    projection_chart_data = simulation[
        [
            "Ano",
            "Subtotal ativo conhecido",
            "Insuficiência total",
            "Total com insuficiência",
            "RCL ajustada projetada",
        ]
    ].copy()
    rcl_series = projection_chart_data["RCL ajustada projetada"].astype(float).mask(
        lambda values: values.eq(0)
    )
    projection_chart_data["Ativos sem aporte"] = (
        projection_chart_data["Subtotal ativo conhecido"].astype(float).div(rcl_series).mul(100)
    )
    projection_chart_data["Insuficiência previdenciária"] = (
        projection_chart_data["Insuficiência total"].astype(float).div(rcl_series).mul(100)
    )
    projection_chart_data["Total com insuficiência"] = (
        projection_chart_data["Total com insuficiência"].astype(float).div(rcl_series).mul(100)
    )
    projection_chart_data = projection_chart_data[
        [
            "Ano",
            "Ativos sem aporte",
            "Insuficiência previdenciária",
            "Total com insuficiência",
        ]
    ]
    projection_chart_long = projection_chart_data.melt(
        id_vars=["Ano"], var_name="Série", value_name="Percentual"
    )
    st.markdown(
        '<div class="fiscal-projection-chart">'
        '<div class="fiscal-projection-chart__title">Trajetória do comprometimento da RCL</div>'
        f'<div class="fiscal-projection-chart__subtitle">Três séries percentuais sobre a RCL ajustada · cenário {html.escape(str(scenario))} · exercício destacado: {int(projection_year)}.</div></div>',
        unsafe_allow_html=True,
    )
    projection_lines = (
        alt.Chart(projection_chart_long)
        .mark_line(point=True, strokeWidth=2.6)
        .encode(
            x=alt.X("Ano:O", title="Exercício", sort=list(range(2027, 2036)), axis=alt.Axis(labelAngle=0)),
            y=alt.Y("Percentual:Q", title="% da RCL ajustada", scale=alt.Scale(zero=True)),
            color=alt.Color(
                "Série:N",
                title=None,
                scale=alt.Scale(
                    domain=["Ativos sem aporte", "Insuficiência previdenciária", "Total com insuficiência"],
                    range=["#17534D", "#A65A46", "#C3A35F"],
                ),
            ),
            tooltip=["Ano:O", "Série:N", alt.Tooltip("Percentual:Q", format=",.2f")],
        )
        .properties(height=285)
    )
    projection_selected = (
        alt.Chart(projection_chart_long.loc[projection_chart_long["Ano"].eq(projection_year)])
        .mark_point(size=135, filled=True, stroke="#FFFFFF", strokeWidth=2)
        .encode(
            x=alt.X("Ano:O", title="Exercício", sort=list(range(2027, 2036))),
            y=alt.Y("Percentual:Q", title="% da RCL ajustada"),
            color=alt.Color(
                "Série:N",
                title=None,
                scale=alt.Scale(
                    domain=["Ativos sem aporte", "Insuficiência previdenciária", "Total com insuficiência"],
                    range=["#17534D", "#A65A46", "#C3A35F"],
                ),
            ),
            tooltip=["Ano:O", "Série:N", alt.Tooltip("Percentual:Q", format=",.2f")],
        )
    )
    st.altair_chart(projection_lines + projection_selected, width="stretch")
    notice(
        "A projeção da DTP oficial, o fechamento de 2026 e qualquer cenário fora do método homologado permanecem bloqueados. "
        "RGF e RREO não são somados e o comprometimento exibido é gerencial, não um indicador oficial da LRF.",
        "risk",
    )
    source_strip(SOURCE_LABELS["fiscal"])
    with st.expander("Memória de cálculo das projeções 2027–2035", expanded=False):
        simulation_view = simulation.copy()
        table_rcl = simulation_view["RCL ajustada projetada"].astype(float).mask(
            lambda values: values.eq(0)
        )
        simulation_view["Comprometimento dos ativos"] = (
            simulation_view["Subtotal ativo conhecido"].astype(float).div(table_rcl).mul(100)
        )
        simulation_view["Comprometimento da insuficiência"] = (
            simulation_view["Insuficiência total"].astype(float).div(table_rcl).mul(100)
        )
        simulation_view["Comprometimento total"] = (
            simulation_view["Total com insuficiência"].astype(float).div(table_rcl).mul(100)
        )
        simulation_view = simulation_view.rename(
            columns={
                "Ano": "Exercício",
                "Cenário-base": "Cenário",
                "Crescimento real anual da folha ativa": "Crescimento salarial",
                "Contribuição patronal civil demonstrável": "Contribuições patronais demonstráveis — civil",
                "Contribuição patronal militar": "Contribuições patronais demonstráveis — militar (N/A)",
            }
        )
        st.dataframe(
            format_table(
                simulation_view,
                money=[
                    "Custo anual da folha ativa",
                    "Contribuições patronais demonstráveis — civil",
                    "Contribuições patronais demonstráveis — militar (N/A)",
                    "Subtotal ativo conhecido",
                    "Insuficiência Financeiro/Civil",
                    "Insuficiência Previdenciário/Civil",
                    "Insuficiência Previdenciário/Militar",
                    "Insuficiência total",
                    "Total com insuficiência",
                    "RCL ajustada projetada",
                ],
                pct=[
                    "Crescimento anual da RCL",
                    "Crescimento salarial",
                    "Comprometimento dos ativos",
                    "Comprometimento da insuficiência",
                    "Comprometimento total",
                    "Comprometimento gerencial",
                ],
            ),
            width="stretch",
            hide_index=True,
        )
        st.caption("A contribuição patronal militar permanece N/A em todos os exercícios; ausência não é zero.")
    notice(
        "Opção A aplicada: o crescimento salarial altera somente as folhas ativas e, por derivação, a contribuição civil sobre sua base. Receitas, despesas, benefícios e insuficiências permanecem nos fluxos originais do DRAA 2026.",
        "info",
    )
    notice(
        "A insuficiência é calculada separadamente para Financeiro/Civil, Previdenciário/Civil e Previdenciário/Militar. Superávit de um plano não compensa déficit de outro.",
        "warn",
    )

    projections = fiscal_projection_scenarios(bundle.fiscal_projections)
    with st.expander("Cenários-base homologados e premissas", expanded=False):
        end_values = projections.loc[projections["Ano"].eq(2035)].sort_values("Cenário")
        st.dataframe(
            format_table(
                end_values[[
                    "Cenário",
                    "Ano",
                    "Crescimento real da folha",
                    "RCL ajustada projetada",
                    "Custo direto estimado",
                    "Comprometimento direto",
                    "Limitação",
                ]],
                money=["RCL ajustada projetada", "Custo direto estimado"],
                pct=["Comprometimento direto"],
            ),
            width="stretch",
            hide_index=True,
        )
        st.caption(
            "Referências diretas parciais preservadas em 2035: A 16,9474%; B 18,7205%; C 20,6588%. "
            "Não são comprometimento ampliado e não precisam coincidir com uma simulação de premissas alteradas."
        )
        st.dataframe(bundle.fiscal_projection_assumptions.copy(), width="stretch", hide_index=True)
        source_strip(SOURCE_LABELS["projection"])
